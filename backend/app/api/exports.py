from io import BytesIO

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.db import get_supabase

router = APIRouter(prefix="/projects/{project_id}/exports", tags=["exports"])


@router.get("/boq.xlsx")
def export_boq_excel(project_id: str):
    client = get_supabase()
    project_result = (
        client.table("projects")
        .select("id,name,location,client_name,currency")
        .eq("id", project_id)
        .limit(1)
        .execute()
    )
    if not project_result.data:
        raise HTTPException(status_code=404, detail="Project not found")

    project = project_result.data[0]
    sections = (
        client.table("boq_sections")
        .select("id,name,sort_order")
        .eq("project_id", project_id)
        .order("sort_order")
        .execute()
    ).data
    items = (
        client.table("boq_items")
        .select("item_no,description,unit,quantity,unit_rate,amount,section_id")
        .eq("project_id", project_id)
        .order("item_no")
        .execute()
    ).data

    section_names = {str(s["id"]): s["name"] for s in sections}

    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ"

    ws["A1"] = project["name"]
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = "Location"
    ws["B2"] = project.get("location") or ""
    ws["A3"] = "Client"
    ws["B3"] = project.get("client_name") or ""
    ws["A4"] = "Currency"
    ws["B4"] = project.get("currency") or "ETB"

    headers = ["Item No.", "Description", "Unit", "Quantity", "Unit Rate", "Amount", "Section"]
    header_row = 6
    for col, header in enumerate(headers, 1):
        cell = ws.cell(header_row, col, header)
        cell.font = Font(bold=True)

    row_num = header_row + 1
    for item in items:
        ws.cell(row_num, 1, item["item_no"])
        ws.cell(row_num, 2, item["description"])
        ws.cell(row_num, 3, item["unit"])
        ws.cell(row_num, 4, float(item["quantity"]))
        ws.cell(row_num, 5, float(item["unit_rate"]))
        ws.cell(row_num, 6, float(item["amount"]))
        section_id = str(item["section_id"]) if item.get("section_id") else None
        ws.cell(row_num, 7, section_names.get(section_id, "Unassigned"))
        row_num += 1

    total_row = row_num + 1
    ws.cell(total_row, 5, "DIRECT COST")
    ws.cell(total_row, 5).font = Font(bold=True)
    ws.cell(total_row, 6, f"=SUM(F{header_row + 1}:F{row_num - 1})")
    ws.cell(total_row, 6).font = Font(bold=True)

    for col in range(1, 8):
        ws.column_dimensions[get_column_letter(col)].width = [12, 42, 12, 14, 16, 18, 24][col - 1]

    for row in ws.iter_rows(min_row=header_row + 1, min_col=4, max_col=6):
        for cell in row:
            cell.number_format = '#,##0.00'

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"boq-{project_id}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
